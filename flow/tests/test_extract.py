# Copyright (c) 2026, Frappe Technologies and contributors
# License: MIT. See LICENSE

from io import BytesIO

from frappe.tests import UnitTestCase

from flow.knowledge.extract import (
	FILE_EXTENSIONS,
	_extract_csv,
	_extract_docx,
	_extract_pptx,
	_looks_like_prose,
	_sniff_extension,
)


class TestDocExtensionSupported(UnitTestCase):
	def test_legacy_doc_extension_is_supported(self):
		self.assertIn("doc", FILE_EXTENSIONS)

	def test_pptx_extension_is_supported(self):
		self.assertIn("pptx", FILE_EXTENSIONS)


class TestExtractDocx(UnitTestCase):
	def _build_docx(self) -> bytes:
		from docx import Document

		document = Document()
		document.add_paragraph("Hello world paragraph.")
		table = document.add_table(rows=2, cols=2)
		table.cell(0, 0).text = "Name"
		table.cell(0, 1).text = "Qty"
		table.cell(1, 0).text = "Widget"
		table.cell(1, 1).text = "5"
		buffer = BytesIO()
		document.save(buffer)
		return buffer.getvalue()

	def test_paragraphs_and_tables_are_both_extracted(self):
		text = _extract_docx(self._build_docx())
		self.assertIn("Hello world paragraph.", text)
		self.assertIn("| Name | Qty |", text)
		self.assertIn("| Widget | 5 |", text)


class TestExtractCsv(UnitTestCase):
	def test_csv_is_structured_as_a_markdown_table(self):
		text = _extract_csv("name,qty\nWidget,5\nGadget,2", delimiter=",")
		self.assertIn("| name | qty |", text)
		self.assertIn("| Widget | 5 |", text)

	def test_oversized_csv_falls_back_to_raw_text(self):
		big = "col\n" + "\n".join(str(i) for i in range(3000))
		text = _extract_csv(big, delimiter=",")
		self.assertEqual(text, big)


class TestExtractPptx(UnitTestCase):
	def test_slide_text_and_tables_are_extracted(self):
		from pptx import Presentation
		from pptx.util import Inches

		presentation = Presentation()
		slide = presentation.slides.add_slide(presentation.slide_layouts[1])
		slide.shapes.title.text = "Title Slide"
		graphic_frame = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(1))
		table = graphic_frame.table
		table.cell(0, 0).text = "A"
		table.cell(0, 1).text = "B"
		table.cell(1, 0).text = "1"
		table.cell(1, 1).text = "2"
		buffer = BytesIO()
		presentation.save(buffer)

		text = _extract_pptx(buffer.getvalue())
		self.assertIn("Title Slide", text)
		self.assertIn("| A | B |", text)


class TestSniffExtension(UnitTestCase):
	def test_sniffs_pdf_magic_bytes(self):
		self.assertEqual(_sniff_extension(b"%PDF-1.4 rest of file"), "pdf")

	def test_sniffs_plain_text(self):
		self.assertEqual(_sniff_extension(b"hello world"), "txt")

	def test_returns_none_for_unrecognized_binary(self):
		self.assertIsNone(_sniff_extension(b"\x00\x01\x02\xff\xfe"))


class TestLooksLikeProse(UnitTestCase):
	def test_accepts_natural_language_sentence(self):
		self.assertTrue(_looks_like_prose("Employee Name: John Smith, Position Title: Manager"))

	def test_rejects_short_run(self):
		self.assertFalse(_looks_like_prose("JFIF"))

	def test_rejects_run_without_word_boundaries(self):
		self.assertFalse(_looks_like_prose("XICC_PROFILEHLinomntrRGB"))

	def test_rejects_symbol_heavy_binary_noise(self):
		self.assertFalse(_looks_like_prose("br0$U$-;DsDm8&7>83mKy"))
